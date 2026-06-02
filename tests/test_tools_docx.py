"""Tests de la tool read_docx.

Si `python-docx` no está instalado, el módulo se skipea con un mensaje claro.
"""

from __future__ import annotations

from pathlib import Path

import pytest

# Skip todo el módulo si python-docx no está disponible
docx_lib = pytest.importorskip("docx")

from wso.tools.docx import read_docx  # noqa: E402


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_simple_docx(path: Path, paragraphs: list[tuple[str, str | None]]) -> Path:
    """Construir un .docx con (texto, style) en orden.

    style=None usa el default. Para heading, pasá "Heading 1" / "Heading 2".
    """
    from docx import Document

    doc = Document()
    for text, style in paragraphs:
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
    doc.save(str(path))
    return path


def _make_docx_with_table(path: Path, rows: list[list[str]]) -> Path:
    """Construir un .docx con una sola tabla."""
    from docx import Document

    doc = Document()
    if not rows:
        doc.save(str(path))
        return path
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r_idx, row_data in enumerate(rows):
        for c_idx, cell_text in enumerate(row_data):
            table.cell(r_idx, c_idx).text = cell_text
    doc.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Path validation
# ---------------------------------------------------------------------------


class TestReadDocxPathValidation:
    def test_relative_path_rejected(self) -> None:
        with pytest.raises(ValueError, match="absoluta"):
            read_docx("relative/file.docx")

    def test_nonexistent_file_raises(self, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            read_docx(str(tmp_path / "no_existe.docx"))

    def test_directory_path_raises(self, tmp_path: Path) -> None:
        with pytest.raises(IsADirectoryError):
            read_docx(str(tmp_path))


# ---------------------------------------------------------------------------
# Reading basic content
# ---------------------------------------------------------------------------


class TestReadDocxBasic:
    def test_returns_string_with_path_header(self, tmp_path: Path) -> None:
        path = tmp_path / "simple.docx"
        _make_simple_docx(path, [("Hola mundo", None)])

        result = read_docx(str(path))
        assert "Archivo:" in result
        assert "Hola mundo" in result

    def test_extracts_multiple_paragraphs(self, tmp_path: Path) -> None:
        path = tmp_path / "multi.docx"
        _make_simple_docx(
            path,
            [("primer parrafo", None), ("segundo parrafo", None), ("tercero", None)],
        )

        result = read_docx(str(path))
        assert "primer parrafo" in result
        assert "segundo parrafo" in result
        assert "tercero" in result

    def test_preserves_paragraph_order(self, tmp_path: Path) -> None:
        path = tmp_path / "order.docx"
        _make_simple_docx(
            path,
            [("alfa", None), ("beta", None), ("gamma", None)],
        )

        result = read_docx(str(path))
        a = result.index("alfa")
        b = result.index("beta")
        g = result.index("gamma")
        assert a < b < g

    def test_skips_empty_paragraphs(self, tmp_path: Path) -> None:
        path = tmp_path / "with_empties.docx"
        _make_simple_docx(
            path,
            [("antes", None), ("", None), ("", None), ("después", None)],
        )

        result = read_docx(str(path))
        assert "antes" in result
        assert "después" in result
        # No debería haber líneas con solo whitespace entre antes/después


# ---------------------------------------------------------------------------
# Headings
# ---------------------------------------------------------------------------


class TestReadDocxHeadings:
    def test_heading_1_gets_single_hash(self, tmp_path: Path) -> None:
        path = tmp_path / "headings.docx"
        _make_simple_docx(
            path,
            [
                ("Título principal", "Heading 1"),
                ("Texto bajo el título", None),
            ],
        )

        result = read_docx(str(path))
        assert "# Título principal" in result

    def test_heading_2_gets_two_hashes(self, tmp_path: Path) -> None:
        path = tmp_path / "h2.docx"
        _make_simple_docx(path, [("Sección", "Heading 2")])

        result = read_docx(str(path))
        assert "## Sección" in result

    def test_heading_3_gets_three_hashes(self, tmp_path: Path) -> None:
        path = tmp_path / "h3.docx"
        _make_simple_docx(path, [("Subsección", "Heading 3")])

        result = read_docx(str(path))
        assert "### Subsección" in result

    def test_normal_paragraph_has_no_prefix(self, tmp_path: Path) -> None:
        path = tmp_path / "normal.docx"
        _make_simple_docx(path, [("Texto normal", None)])

        result = read_docx(str(path))
        # No debería tener prefijo #
        line = next((ln for ln in result.splitlines() if "Texto normal" in ln), "")
        assert not line.lstrip().startswith("#")


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


class TestReadDocxTables:
    def test_table_serialized_with_pipe(self, tmp_path: Path) -> None:
        path = tmp_path / "table.docx"
        _make_docx_with_table(
            path,
            [
                ["Header1", "Header2", "Header3"],
                ["a1", "b1", "c1"],
                ["a2", "b2", "c2"],
            ],
        )

        result = read_docx(str(path))
        assert "| Header1 | Header2 | Header3 |" in result
        assert "| a1 | b1 | c1 |" in result
        assert "| a2 | b2 | c2 |" in result

    def test_pipe_in_cell_content_is_escaped(self, tmp_path: Path) -> None:
        path = tmp_path / "pipe.docx"
        _make_docx_with_table(path, [["text | with pipe", "ok"]])

        result = read_docx(str(path))
        # El pipe interno debe estar escapado para no confundir con separador
        assert "\\|" in result


# ---------------------------------------------------------------------------
# Corrupt files
# ---------------------------------------------------------------------------


class TestReadDocxCorrupt:
    def test_non_docx_file_raises_value_error(self, tmp_path: Path) -> None:
        bad = tmp_path / "not_a_docx.docx"
        bad.write_text("this is not a DOCX, just plain text")

        with pytest.raises(ValueError, match="abrir"):
            read_docx(str(bad))


# ---------------------------------------------------------------------------
# Registry integration
# ---------------------------------------------------------------------------


class TestReadDocxRegistration:
    def test_registered_in_builtin_tools(self) -> None:
        from wso.tools.registry import load_builtin_tools

        registry = load_builtin_tools()
        assert "read_docx" in registry

        defn = registry.get("read_docx")
        assert defn is not None
        assert defn.category.value == "read"

    def test_appears_in_prompt_section(self) -> None:
        from wso.tools.registry import load_builtin_tools

        registry = load_builtin_tools()
        section = registry.to_prompt_section()
        assert "## read_docx" in section


class TestReadDocxMaxChars:
    """max_chars: acotar el tamaño de la lectura para no desbordar contexto."""

    def test_long_docx_truncated_with_notice(self, tmp_path: Path) -> None:
        paras = [("Párrafo " + str(i) + " " + "Z" * 400, None) for i in range(40)]
        docx_path = _make_simple_docx(tmp_path / "long.docx", paras)
        out = read_docx(str(docx_path), max_chars=2000)
        assert "TRUNCADO" in out
        assert "max_chars" in out
        assert len(out) < 3000

    def test_short_docx_not_truncated(self, tmp_path: Path) -> None:
        docx_path = _make_simple_docx(
            tmp_path / "short.docx", [("Hola mundo", None)]
        )
        out = read_docx(str(docx_path), max_chars=10000)
        assert "TRUNCADO" not in out
        assert "Hola mundo" in out
