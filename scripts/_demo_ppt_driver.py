"""Driver de demo: arma un PPT free-form pasando código python-pptx por run_python.

NO es parte del harness — es solo para demostrar que run_python puede construir
una presentación libre (sin template, sin el schema de generate_pptx).
"""
from __future__ import annotations

from pathlib import Path

from wso.config import settings
from wso.tools.code import run_python

OUT = settings.workspace_dir / "output" / "demo_ppt_libre.pptx"

PPTX_CODE = f'''
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paleta del estudio
INK   = RGBColor(0x14, 0x14, 0x16)   # casi negro
PAPER = RGBColor(0xF5, 0xF3, 0xEE)   # blanco hueso
ACC   = RGBColor(0xC8, 0x55, 0x3D)   # terracota
GREY  = RGBColor(0x8A, 0x8A, 0x8A)

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]       # layout en blanco = control total

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def box(slide, l, t, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT, font="Arial"):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
    r.font.color.rgb = color
    return tb

def rect(slide, l, t, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background()
    return sp

# --- Slide 1: portada full-bleed ---
s = prs.slides.add_slide(BLANK); bg(s, INK)
rect(s, 0.8, 3.05, 1.6, 0.10, ACC)                       # línea de acento
box(s, 0.8, 3.2, 11, 1.4, "White Suit Studio", 54, PAPER, bold=True)
box(s, 0.82, 4.5, 11, 0.8, "Capacidades de presentación libre — armado con python-pptx", 20, GREY)
box(s, 0.82, 6.7, 11, 0.5, "Demo · run_python", 12, ACC)

# --- Slide 2: intro con bullets custom ---
s = prs.slides.add_slide(BLANK); bg(s, PAPER)
box(s, 0.8, 0.7, 11, 1, "Por qué libre", 36, INK, bold=True)
rect(s, 0.85, 1.7, 1.2, 0.07, ACC)
bullets = [
    "Sin template: control total del layout, color y tipografía.",
    "Formas, tablas, imágenes y charts arbitrarios.",
    "El modelo escribe el código; run_python lo ejecuta contenido.",
]
top = 2.2
for b in bullets:
    rect(s, 0.9, top + 0.12, 0.18, 0.18, ACC, MSO_SHAPE.OVAL)
    box(s, 1.3, top, 10.5, 0.7, b, 20, INK)
    top += 0.95

# --- Slide 3: dos tarjetas ---
s = prs.slides.add_slide(BLANK); bg(s, PAPER)
box(s, 0.8, 0.7, 11, 1, "Dos caminos", 36, INK, bold=True)
rect(s, 0.85, 1.7, 1.2, 0.07, ACC)
def card(l, titulo, cuerpo, color):
    rect(s, l, 2.3, 5.4, 3.6, color, MSO_SHAPE.ROUNDED_RECTANGLE)
    box(s, l + 0.4, 2.7, 4.6, 0.8, titulo, 24, PAPER, bold=True)
    box(s, l + 0.4, 3.6, 4.6, 2, cuerpo, 16, PAPER)
card(0.9, "generate_pptx", "Estructura JSON: títulos y bullets, sin template. Rápido y predecible.", INK)
card(6.9, "run_python", "Código python-pptx directo: cualquier diseño. Máxima libertad.", ACC)

# --- Slide 4: tabla ---
s = prs.slides.add_slide(BLANK); bg(s, PAPER)
box(s, 0.8, 0.7, 11, 1, "Servicios", 36, INK, bold=True)
rect(s, 0.85, 1.7, 1.2, 0.07, ACC)
rows, cols = 4, 3
gt = s.shapes.add_table(rows, cols, Inches(0.9), Inches(2.3), Inches(11.5), Inches(3)).table
data = [["Servicio", "Entrega", "Desde"],
        ["Identidad de marca", "4 semanas", "USD 3.500"],
        ["Sitio web", "6 semanas", "USD 5.000"],
        ["Deck de inversión", "2 semanas", "USD 1.800"]]
for r_i in range(rows):
    for c_i in range(cols):
        cell = gt.cell(r_i, c_i)
        cell.text = data[r_i][c_i]
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(16)
        para.font.bold = (r_i == 0)
        para.font.color.rgb = PAPER if r_i == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK if r_i == 0 else PAPER

# --- Slide 5: cierre ---
s = prs.slides.add_slide(BLANK); bg(s, INK)
box(s, 0, 3.0, 13.333, 1.2, "Hagamos algo lindo.", 44, PAPER, bold=True, align=PP_ALIGN.CENTER)
box(s, 0, 4.3, 13.333, 0.6, "hola@whitesuit.studio", 18, ACC, align=PP_ALIGN.CENTER)

prs.save(r"{OUT}")
print("Slides:", len(prs.slides._sldIdLst), "→", r"{OUT}")
'''


def main() -> int:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    print(run_python(PPTX_CODE, timeout_s=60))
    print("existe:", OUT.exists(), "| bytes:", OUT.stat().st_size if OUT.exists() else 0)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
