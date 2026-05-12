"""Tools de PowerPoint (.pptx) basadas en python-pptx.

Cuatro tools declarativas para que el agente pueda:

    - generate_pptx               : crear un .pptx desde una estructura JSON
    - read_pptx                   : extraer el texto (títulos, bullets, notas) de un .pptx
    - edit_pptx_slide             : modificar un slide individual sin recrear el deck
    - generate_pptx_from_template : crear un .pptx reutilizando layouts/branding de un template

Convención de la lista de slides
--------------------------------

El modelo emite la estructura de slides como JSON dentro de un solo arg
string (`slides_json` o `updates_json`). Esto es un workaround a la
limitación del parser XML (no soporta XML anidado — ver decisión 3.15
en DESIGN.md).

Cada elemento de la lista valida contra el discriminator union `Slide`
definido en `pptx_schemas.py`. El campo `layout` decide la forma:

    [
      {"layout": "title",   "title": "Q3 Review", "subtitle": "2025"},
      {"layout": "content", "title": "Highlights", "bullets": ["...", "..."]},
      {"layout": "two_content", "title": "Cliente vs Estudio",
         "left_bullets": ["..."], "right_bullets": ["..."]},
      {"layout": "image",   "title": "Brand board", "image_path": "/abs/path.png"},
      {"layout": "section_header", "title": "Anexos"},
      {"layout": "blank"}
    ]

Errores
-------

Como las demás tools, las excepciones se propagan crudas. El loop las
captura y las formatea como observation para que el modelo reintente:

    - ValueError              : path relativo / slides_json inválido
    - FileNotFoundError       : archivo / imagen no existe
    - IsADirectoryError       : el path apunta a un directorio
    - json.JSONDecodeError    : slides_json no es JSON válido
    - pydantic.ValidationError: la estructura no matchea ningún layout
    - IndexError              : slide_index fuera de rango (edit_pptx_slide)
    - ImportError             : python-pptx no está instalado (mensaje guía)
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pydantic import ValidationError

from wso.tools.base import PermissionCategory, tool
from wso.tools.pptx_schemas import (
    LAYOUT_INDEX,
    BlankSlide,
    ContentSlide,
    ImageSlide,
    SectionHeaderSlide,
    SlideDeck,
    TitleSlide,
    TwoContentSlide,
)

# ---------------------------------------------------------------------------
# Lazy import de python-pptx
# ---------------------------------------------------------------------------


def _require_pptx() -> Any:
    """Importar python-pptx con un mensaje claro si falta la dep.

    No importamos a nivel módulo para que `wso` siga corriendo aunque
    la extra `[office]` no esté instalada. Solo falla si el modelo
    realmente invoca una tool de pptx.
    """
    try:
        import pptx  # noqa: PLC0415

        return pptx
    except ImportError as e:  # pragma: no cover — depende del entorno
        raise ImportError(
            "python-pptx no está instalado. Para usar las tools de PPTX, "
            "instalá la extra: `pip install -e \".[office]\"` "
            "o `pip install python-pptx`."
        ) from e


# ---------------------------------------------------------------------------
# Helpers de path
# ---------------------------------------------------------------------------


def _validate_absolute(path: str, label: str = "ruta") -> Path:
    """Expandir `~` y verificar que el path sea absoluto."""
    p = Path(path).expanduser()
    if not p.is_absolute():
        raise ValueError(
            f"La {label} debe ser absoluta: {path!r}. "
            f"Usá rutas tipo '/Users/...' o '~/...'."
        )
    return p


def _validate_pptx_path_for_read(path: str) -> Path:
    """Validar un path existente que apunta a un .pptx legible."""
    p = _validate_absolute(path)
    if not p.exists():
        raise FileNotFoundError(f"El .pptx no existe: {p}")
    if p.is_dir():
        raise IsADirectoryError(
            f"La ruta es un directorio, no un archivo .pptx: {p}"
        )
    return p


# ---------------------------------------------------------------------------
# Helpers de aplicación de slides (rellenan placeholders de un layout)
# ---------------------------------------------------------------------------


def _set_text_frame_bullets(text_frame: Any, bullets: list[str]) -> None:
    """Reemplazar el contenido de un text_frame con una lista de bullets.

    python-pptx no tiene API para "clear & set bullets" en un placeholder,
    así que el primer bullet va al párrafo existente (que viene vacío por
    default en un slide nuevo) y los siguientes se agregan.
    """
    if not bullets:
        text_frame.text = ""
        return
    text_frame.text = bullets[0]
    for bullet in bullets[1:]:
        p = text_frame.add_paragraph()
        p.text = bullet


def _set_speaker_notes(slide: Any, notes: str) -> None:
    """Setear el contenido de las speaker notes de un slide."""
    if not notes:
        return
    slide.notes_slide.notes_text_frame.text = notes


def _apply_title_slide(slide: Any, data: TitleSlide) -> None:
    """Llenar un layout 'title': título principal + subtítulo."""
    slide.shapes.title.text = data.title
    # placeholder 1 = subtitle en el layout 0 (Title Slide)
    if data.subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = data.subtitle
    _set_speaker_notes(slide, data.notes)


def _apply_content_slide(slide: Any, data: ContentSlide) -> None:
    """Llenar un layout 'content': título + bullets."""
    slide.shapes.title.text = data.title
    # placeholder 1 = body/content en el layout 1 (Title and Content)
    if len(slide.placeholders) > 1:
        _set_text_frame_bullets(slide.placeholders[1].text_frame, data.bullets)
    _set_speaker_notes(slide, data.notes)


def _apply_section_header_slide(slide: Any, data: SectionHeaderSlide) -> None:
    """Llenar un layout 'section_header'."""
    slide.shapes.title.text = data.title
    if data.subtitle and len(slide.placeholders) > 1:
        slide.placeholders[1].text = data.subtitle
    _set_speaker_notes(slide, data.notes)


def _apply_two_content_slide(slide: Any, data: TwoContentSlide) -> None:
    """Llenar un layout 'two_content': título + dos columnas de bullets.

    En el layout 3 (Two Content) los placeholders típicamente son:
        - 0: title
        - 1: left content
        - 2: right content
    """
    slide.shapes.title.text = data.title
    placeholders = list(slide.placeholders)
    # Filtrar el title placeholder y quedarnos con los content
    content_phs = [ph for ph in placeholders if ph != slide.shapes.title]
    if len(content_phs) >= 1:
        _set_text_frame_bullets(content_phs[0].text_frame, data.left_bullets)
    if len(content_phs) >= 2:
        _set_text_frame_bullets(content_phs[1].text_frame, data.right_bullets)
    _set_speaker_notes(slide, data.notes)


def _apply_image_slide(slide: Any, data: ImageSlide) -> None:
    """Llenar un layout 'image': título + imagen embebida.

    La imagen se ancla a una posición fija (centrada, ~80% del slide).
    Si el caption está presente, se inserta como speaker note o como
    text box separado — por simplicidad, va a notas.
    """
    pptx = _require_pptx()
    Inches = pptx.util.Inches  # noqa: N806

    image_path = _validate_absolute(data.image_path, label="image_path")
    if not image_path.exists():
        raise FileNotFoundError(
            f"La imagen no existe: {image_path} (referenciada en slide image)"
        )

    if slide.shapes.title is not None:
        slide.shapes.title.text = data.title

    # Insertar imagen: top-left aprox al 10% / 25% del slide,
    # con ancho ~80% (slide estándar 10"×7.5", pptx default).
    slide.shapes.add_picture(
        str(image_path),
        left=Inches(1),
        top=Inches(1.75),
        width=Inches(8),
    )

    notes = data.notes
    if data.caption:
        notes = f"{data.caption}\n\n{notes}".strip() if notes else data.caption
    _set_speaker_notes(slide, notes)


def _apply_blank_slide(slide: Any, data: BlankSlide) -> None:
    """Llenar un layout 'blank' — solo título opcional."""
    if data.title and slide.shapes.title is not None:
        slide.shapes.title.text = data.title
    _set_speaker_notes(slide, data.notes)


def _apply_slide(prs: Any, slide_data: Any) -> None:
    """Agregar un slide al deck según el tipo de slide_data."""
    layout_name = slide_data.layout
    layout_idx = LAYOUT_INDEX[layout_name]
    if layout_idx >= len(prs.slide_layouts):
        raise IndexError(
            f"El layout {layout_name!r} (índice {layout_idx}) no existe en este deck. "
            f"El template tiene {len(prs.slide_layouts)} layouts. "
            f"Probá con un template distinto o un layout más simple."
        )

    layout = prs.slide_layouts[layout_idx]
    slide = prs.slides.add_slide(layout)

    if isinstance(slide_data, TitleSlide):
        _apply_title_slide(slide, slide_data)
    elif isinstance(slide_data, ContentSlide):
        _apply_content_slide(slide, slide_data)
    elif isinstance(slide_data, SectionHeaderSlide):
        _apply_section_header_slide(slide, slide_data)
    elif isinstance(slide_data, TwoContentSlide):
        _apply_two_content_slide(slide, slide_data)
    elif isinstance(slide_data, ImageSlide):
        _apply_image_slide(slide, slide_data)
    elif isinstance(slide_data, BlankSlide):
        _apply_blank_slide(slide, slide_data)
    else:  # pragma: no cover — guard defensivo
        raise TypeError(f"Tipo de slide no soportado: {type(slide_data).__name__}")


def _parse_slides_json(slides_json: str) -> SlideDeck:
    """Parsear y validar un JSON string contra el schema SlideDeck.

    Acepta tanto un array JSON top-level (caso típico que emite el modelo)
    como un objeto `{"slides": [...]}`.

    Errores se propagan: json.JSONDecodeError o pydantic.ValidationError.
    """
    raw = json.loads(slides_json)
    if isinstance(raw, list):
        return SlideDeck(slides=raw)
    if isinstance(raw, dict) and "slides" in raw:
        return SlideDeck.model_validate(raw)
    raise ValueError(
        "slides_json debe ser un array JSON de slides o un objeto "
        '{"slides": [...]}, recibido: '
        f"{type(raw).__name__}"
    )


def _delete_all_slides(prs: Any) -> None:
    """Eliminar todos los slides de un Presentation, preservando layouts/master.

    python-pptx no expone una API limpia para esto. La manera correcta es:
      1. Recolectar los rIds (relationship IDs) de cada sldId.
      2. Quitar los sldId del sldIdLst (saca la referencia en el XML).
      3. Soltar las relaciones (drop_rel) para que las partes XML asociadas
         no queden colgadas en el package — si no, al agregar slides nuevos
         se generan colisiones de nombre en el .zip y el writer escupe
         UserWarning de duplicate name.

    Se usa para que `generate_pptx_from_template` empiece con un deck vacío
    que conserva el branding del template (master + layouts).
    """
    sld_id_lst = prs.slides._sldIdLst
    r_ids = [sld_id.rId for sld_id in list(sld_id_lst)]
    for sld_id in list(sld_id_lst):
        sld_id_lst.remove(sld_id)
    for r_id in r_ids:
        prs.part.drop_rel(r_id)


# ---------------------------------------------------------------------------
# Tool: generate_pptx
# ---------------------------------------------------------------------------


@tool(
    name="generate_pptx",
    category=PermissionCategory.WRITE,
    description=(
        "Genera un archivo .pptx desde una estructura JSON de slides. "
        "Cada slide tiene un campo 'layout' que decide su forma: "
        "title, content, section_header, two_content, image, blank. "
        "Soporta speaker notes y bullets. "
        "Para imágenes, image_path debe ser ruta local absoluta."
    ),
    args_schema={
        "output_path": "ruta absoluta donde escribir el .pptx (ej: /Users/coco/.../deck.pptx)",
        "slides_json": (
            "JSON array con los slides del deck. Ej: "
            '[{"layout":"title","title":"Hola","subtitle":"2026"},'
            '{"layout":"content","title":"Puntos","bullets":["a","b"]}]'
        ),
    },
)
def generate_pptx(output_path: str, slides_json: str) -> str:
    """Crear un .pptx nuevo desde la estructura declarativa."""
    pptx = _require_pptx()

    target = _validate_absolute(output_path, label="output_path")
    if target.is_dir():
        raise IsADirectoryError(
            f"output_path es un directorio: {target}. "
            f"Pasá una ruta a un archivo .pptx."
        )

    deck = _parse_slides_json(slides_json)
    target.parent.mkdir(parents=True, exist_ok=True)

    prs = pptx.Presentation()
    for slide_data in deck.slides:
        _apply_slide(prs, slide_data)

    prs.save(str(target))
    return (
        f"Generado: {target} "
        f"({len(deck.slides)} slide{'s' if len(deck.slides) != 1 else ''})"
    )


# ---------------------------------------------------------------------------
# Tool: read_pptx
# ---------------------------------------------------------------------------


@tool(
    name="read_pptx",
    category=PermissionCategory.READ,
    description=(
        "Extrae el contenido textual de un archivo .pptx existente: "
        "títulos, bullets, texto en text boxes y speaker notes. "
        "Devuelve un string estructurado por slide. "
        "Útil para revisar un deck antes de modificarlo o resumirlo."
    ),
    args_schema={
        "path": "ruta absoluta del archivo .pptx a leer",
    },
)
def read_pptx(path: str) -> str:
    """Extraer texto de un .pptx, slide por slide."""
    pptx = _require_pptx()

    p = _validate_pptx_path_for_read(path)
    prs = pptx.Presentation(str(p))

    lines: list[str] = [f"Archivo: {p}", f"Slides: {len(prs.slides)}", ""]
    for idx, slide in enumerate(prs.slides):
        lines.append(f"--- Slide {idx} ---")

        # Título (si hay placeholder de title)
        title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title_text = slide.shapes.title.text_frame.text.strip()
            if title_text:
                title = title_text
                lines.append(f"Título: {title}")

        # Resto del texto: iteramos shapes con text_frame, excluyendo el title
        body_chunks: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if text:
                body_chunks.append(text)
        if body_chunks:
            lines.append("Contenido:")
            for chunk in body_chunks:
                # Indentar cada línea del chunk
                for body_line in chunk.splitlines():
                    lines.append(f"  • {body_line}")

        # Speaker notes
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                lines.append("Notas:")
                for note_line in notes_text.splitlines():
                    lines.append(f"  {note_line}")

        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Tool: edit_pptx_slide
# ---------------------------------------------------------------------------


@tool(
    name="edit_pptx_slide",
    category=PermissionCategory.WRITE,
    description=(
        "Modifica un slide individual de un .pptx existente sin recrear el deck. "
        "Solo cambia los campos provistos en updates_json (title, bullets, notes). "
        "Preserva el resto del deck. Útil para correcciones puntuales."
    ),
    args_schema={
        "path": "ruta absoluta del archivo .pptx a editar (se modifica in-place)",
        "slide_index": "índice del slide a editar, base 0 (primer slide = 0)",
        "updates_json": (
            "JSON con los campos a actualizar. Claves soportadas: "
            '"title" (str), "bullets" (list[str], reemplaza el body), '
            '"notes" (str, reemplaza speaker notes). '
            'Ej: {"title": "Nuevo título", "bullets": ["a", "b"]}'
        ),
    },
)
def edit_pptx_slide(path: str, slide_index: int, updates_json: str) -> str:
    """Modificar un slide existente."""
    pptx = _require_pptx()

    p = _validate_pptx_path_for_read(path)

    try:
        updates = json.loads(updates_json)
    except json.JSONDecodeError as e:
        raise ValueError(f"updates_json no es JSON válido: {e}") from e

    if not isinstance(updates, dict):
        raise ValueError(
            f"updates_json debe ser un objeto JSON, recibido: {type(updates).__name__}"
        )

    allowed_keys = {"title", "bullets", "notes"}
    extra = set(updates) - allowed_keys
    if extra:
        raise ValueError(
            f"Claves no soportadas en updates_json: {sorted(extra)}. "
            f"Permitidas: {sorted(allowed_keys)}."
        )

    prs = pptx.Presentation(str(p))

    if slide_index < 0 or slide_index >= len(prs.slides):
        raise IndexError(
            f"slide_index fuera de rango: {slide_index}. "
            f"El deck tiene {len(prs.slides)} slides (índices 0..{len(prs.slides) - 1})."
        )

    slide = prs.slides[slide_index]
    changes: list[str] = []

    if "title" in updates:
        new_title = updates["title"]
        if not isinstance(new_title, str):
            raise ValueError(f"title debe ser str, recibido: {type(new_title).__name__}")
        if slide.shapes.title is None:
            raise ValueError(
                f"Slide {slide_index} no tiene placeholder de título. "
                f"No se puede actualizar 'title' en este layout."
            )
        slide.shapes.title.text = new_title
        changes.append("title")

    if "bullets" in updates:
        new_bullets = updates["bullets"]
        if not isinstance(new_bullets, list) or not all(
            isinstance(b, str) for b in new_bullets
        ):
            raise ValueError("bullets debe ser una lista de strings.")
        body_ph = next(
            (ph for ph in slide.placeholders if ph != slide.shapes.title and ph.has_text_frame),
            None,
        )
        if body_ph is None:
            raise ValueError(
                f"Slide {slide_index} no tiene placeholder de contenido. "
                f"No se pueden actualizar 'bullets' en este layout."
            )
        _set_text_frame_bullets(body_ph.text_frame, new_bullets)
        changes.append(f"bullets ({len(new_bullets)})")

    if "notes" in updates:
        new_notes = updates["notes"]
        if not isinstance(new_notes, str):
            raise ValueError(f"notes debe ser str, recibido: {type(new_notes).__name__}")
        slide.notes_slide.notes_text_frame.text = new_notes
        changes.append("notes")

    if not changes:
        return f"Sin cambios: updates_json no contenía claves actualizables. ({p})"

    prs.save(str(p))
    return f"Slide {slide_index} actualizado ({', '.join(changes)}) en {p}"


# ---------------------------------------------------------------------------
# Tool: generate_pptx_from_template
# ---------------------------------------------------------------------------


@tool(
    name="generate_pptx_from_template",
    category=PermissionCategory.WRITE,
    description=(
        "Genera un .pptx usando un template existente como base (preserva "
        "branding, colores, fuentes y layouts del template). "
        "Elimina los slides del template y agrega los nuevos definidos en slides_json. "
        "Útil para usar el template corporativo de White Suit."
    ),
    args_schema={
        "template_path": "ruta absoluta al .pptx que se usa como template",
        "output_path": "ruta absoluta donde escribir el .pptx generado",
        "slides_json": (
            "JSON array con los slides nuevos a crear. Mismo formato que generate_pptx. "
            'Ej: [{"layout":"title","title":"Hola"}]'
        ),
    },
)
def generate_pptx_from_template(
    template_path: str, output_path: str, slides_json: str
) -> str:
    """Generar un .pptx reutilizando los layouts/branding de un template."""
    pptx = _require_pptx()

    template = _validate_pptx_path_for_read(template_path)
    target = _validate_absolute(output_path, label="output_path")
    if target.is_dir():
        raise IsADirectoryError(
            f"output_path es un directorio: {target}. "
            f"Pasá una ruta a un archivo .pptx."
        )

    deck = _parse_slides_json(slides_json)
    target.parent.mkdir(parents=True, exist_ok=True)

    prs = pptx.Presentation(str(template))
    _delete_all_slides(prs)

    for slide_data in deck.slides:
        _apply_slide(prs, slide_data)

    prs.save(str(target))
    return (
        f"Generado: {target} desde template {template.name} "
        f"({len(deck.slides)} slide{'s' if len(deck.slides) != 1 else ''})"
    )


# ---------------------------------------------------------------------------
# Sentinela para silenciar pyright/mypy "unused import" en consumidores
# (los modelos de pydantic se importan solo para validación, no se usan
# directamente acá)
# ---------------------------------------------------------------------------


_VALIDATION_SUPPORT = ValidationError
